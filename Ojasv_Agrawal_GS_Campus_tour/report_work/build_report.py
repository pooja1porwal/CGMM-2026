from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "report_work" / "source"
OUT = ROOT / "SGSITS-Virtual-Campus-Project-Report.docx"
NAVY = "0B2545"
GOLD = "BFA15F"
PALE_GOLD = "F5F0E6"
SLATE = "64748B"
LIGHT = "F7F8FA"
WHITE = "FFFFFF"
BLACK = "172033"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_inches):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = [int(x * 1440) for x in widths_inches]
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[i]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_inches[i])
            set_cell_margins(cell)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def font(run, name="Arial", size=None, color=BLACK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, text, end):
        run._r.append(el)
    font(run, size=9, color=SLATE)


def set_picture_alt(run, description):
    nodes = run._r.xpath(".//wp:docPr")
    if nodes:
        nodes[0].set("descr", description)
        nodes[0].set("title", description)


def paragraph_border_bottom(paragraph, color=GOLD, size=10, space=5):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def shade_paragraph(paragraph, fill=LIGHT):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_para(doc, text="", style=None, bold_prefix=None, align=None, before=0, after=7):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        font(r)
    else:
        r = p.add_run(text)
        font(r)
    return p


def add_callout(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.12
    shade_paragraph(p, PALE_GOLD)
    r = p.add_run(f"{label.upper()}  ")
    font(r, size=9.5, color=GOLD, bold=True)
    r = p.add_run(text)
    font(r, size=10, color=NAVY)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        font(p.add_run(item), size=10.5)


def add_numbers(doc, items):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); lvl.append(start)
    fmt = OxmlElement("w:numFmt"); fmt.set(qn("w:val"), "decimal"); lvl.append(fmt)
    text = OxmlElement("w:lvlText"); text.set(qn("w:val"), "%1."); lvl.append(text)
    suff = OxmlElement("w:suff"); suff.set(qn("w:val"), "tab"); lvl.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), "360"); tabs.append(tab); p_pr.append(tabs)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "360"); ind.set(qn("w:hanging"), "240"); p_pr.append(ind)
    lvl.append(p_pr); abstract.append(lvl); numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    abs_id = OxmlElement("w:abstractNumId"); abs_id.set(qn("w:val"), str(abstract_id)); num.append(abs_id); numbering.append(num)
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        pPr = p._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
        numId = OxmlElement("w:numId"); numId.set(qn("w:val"), str(num_id))
        numPr.append(ilvl); numPr.append(numId); pPr.append(numPr)
        font(p.add_run(item), size=10.5)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for i, value in enumerate(headers):
        set_cell_shading(hdr.cells[i], NAVY)
        hdr.cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr.cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        font(r, size=9.5, color=WHITE, bold=True)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        if ri % 2:
            for c in cells:
                set_cell_shading(c, LIGHT)
        for i, value in enumerate(row):
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            font(r, size=9.25, color=BLACK, bold=(i == 0))
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    shade_paragraph(p, "EEF2F6")
    for i, line in enumerate(lines):
        r = p.add_run(line + ("\n" if i < len(lines) - 1 else ""))
        font(r, name="Consolas", size=8.5, color=NAVY)


def add_figure(doc, number, filename, caption, width=6.85):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(SRC / filename), width=Inches(width))
    set_picture_alt(run, caption)
    cap = doc.add_paragraph()
    cap.style = "Caption"
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3)
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(f"Figure {number}. {caption}")
    font(r, size=9, color=SLATE, italic=True)


def add_page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.35)

# Styles
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.15

for name, size, color, before, after in (
    ("Heading 1", 18, NAVY, 14, 8),
    ("Heading 2", 13, NAVY, 10, 5),
    ("Heading 3", 11, GOLD, 8, 4),
):
    st = doc.styles[name]
    st.font.name = "Cambria"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Number"):
    st = doc.styles[name]
    st.font.name = "Arial"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    st.font.size = Pt(10.5)
    st.paragraph_format.left_indent = Inches(0.25)
    st.paragraph_format.first_line_indent = Inches(-0.18)

caption = doc.styles["Caption"]
caption.font.name = "Arial"
caption.font.size = Pt(9)
caption.font.italic = True
caption.font.color.rgb = RGBColor.from_string(SLATE)

# Running header/footer
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("SGSITS VIRTUAL CAMPUS  |  COMPUTER GRAPHICS PROJECT")
font(hr, size=8.5, color=SLATE, bold=True)
paragraph_border_bottom(hp, color=GOLD, size=6, space=4)
footer = section.footer
fp = footer.paragraphs[0]
fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.7), WD_TAB_ALIGNMENT.RIGHT)
font(fp.add_run("Ojasv Agrawal  |  SGSITS Indore"), size=8.5, color=SLATE)
font(fp.add_run("\tPage "), size=8.5, color=SLATE)
add_field(fp, "PAGE")

# Cover page
for _ in range(2):
    doc.add_paragraph()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
logo_run = p.add_run()
logo_run.add_picture(str(ROOT / "apps" / "frontend" / "public" / "sgsits-logo.png"), width=Inches(1.35))
set_picture_alt(logo_run, "SGSITS institute emblem")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(3)
r = p.add_run("SHRI G. S. INSTITUTE OF TECHNOLOGY & SCIENCE, INDORE")
font(r, name="Cambria", size=15, color=NAVY, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(24)
r = p.add_run("Department of Computer Engineering  |  Computer Graphics Project")
font(r, size=10, color=SLATE, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run("PROJECT REPORT")
font(r, size=10, color=GOLD, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("SGSITS Virtual Campus")
font(r, name="Cambria", size=30, color=NAVY, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(26)
r = p.add_run("An Interactive Map-Based 360 Degree College Tour and Computer Graphics Navigation System")
font(r, name="Cambria", size=14, color=SLATE, italic=True)
paragraph_border_bottom(p, color=GOLD, size=10, space=8)

add_table(doc, ["Submitted by", "Submitted to"], [["Ojasv Agrawal", "Dr. Puja Gupta"]], [3.37, 3.4])
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(30)
font(p.add_run("Academic Session 2026-27"), size=11, color=NAVY, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("Live Project: https://mp-darshan-360.vercel.app/"), size=9.5, color=SLATE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(p.add_run("Repository: https://github.com/ojasvagr123/mp-darshan-360"), size=9.5, color=SLATE)

add_page_break(doc)

# Front matter
h = doc.add_heading("Project Summary", level=1)
paragraph_border_bottom(h)
add_callout(doc, "Project overview", "SGSITS Virtual Campus is a full-stack interactive college-tour application. A public spectator can explore institute landmarks from a campus image map or a sliding place directory, follow an animated route, open a 360 degree panorama, read place information, browse photographs and listen to an audio guide. A single college administrator manages tour stops by clicking the campus map and uploading content.")
add_table(doc, ["Field", "Details"], [
    ("Project Name", "SGSITS Virtual Campus"),
    ("Domain", "Education, campus navigation, multimedia guide"),
    ("Problem Statement", "Interactive College Tour Guide"),
    ("Public User", "Spectator / visitor / prospective student"),
    ("Privileged User", "College administrator only"),
    ("Frontend", "React, Vite, HTML Canvas"),
    ("Backend", "Node.js, Express.js, Prisma ORM"),
    ("Database", "SQLite locally; Supabase PostgreSQL in production"),
    ("Deployment", "Vercel frontend, Render backend, Supabase database"),
    ("Submitted By", "Ojasv Agrawal"),
    ("Submitted To", "Dr. Puja Gupta"),
], [1.65, 5.12])

doc.add_heading("Abstract", level=1)
add_para(doc, "College campuses can be difficult to understand for prospective students, parents, guests and new learners who cannot physically visit beforehand. Conventional maps provide positions but do not communicate what a place looks like, why it matters, or how it connects to other locations. This project transforms the original MP tourism concept into SGSITS Virtual Campus, an institute-specific multimedia tour built around an image map of Shri G. S. Institute of Technology & Science, Indore.")
add_para(doc, "The application supports two user modes. The spectator requires no account and can select a place directly from the campus map or from a searchable sliding directory. The administrator signs in through a protected account, clicks an exact map point, enters descriptions and accessibility details, and uploads a panorama and gallery images. The frontend displays the published stop through 360 view, about, gallery, audio guide and CG navigation tabs.")
add_para(doc, "Computer Graphics is integrated into the application rather than presented as an isolated demonstration. Normalized image coordinates provide viewport mapping; the DDA algorithm rasterizes a route from the Main Gate; a quadratic Bezier curve controls animated movement; translation and rotation orient the direction marker; and requestAnimationFrame continuously redraws the Canvas. The result is a deployable full-stack project that connects graphics theory with a useful campus-discovery problem.")

doc.add_heading("Table of Contents", level=1)
toc = [
    "Problem Statement", "Proposed Solution", "Objectives", "Scope of the Project",
    "Technology Stack", "System Architecture", "User Roles and Main Flow",
    "Major Application Modules", "Database Design", "API Design",
    "Backend Processing Logic", "Computer Graphics Implementation",
    "Frontend UI/UX and Application Screens", "Local Setup Guide",
    "Deployment Architecture", "Testing and Validation", "Security Considerations",
    "Limitations", "Future Scope", "Conclusion", "References"
]
add_numbers(doc, toc)

add_page_break(doc)

# Core report
h = doc.add_heading("1. Problem Statement", level=1); paragraph_border_bottom(h)
add_para(doc, "The revised project problem is to design an Interactive College Tour Guide for SGSITS Indore. A visitor should be able to understand important campus locations without requiring physical presence, specialist software or a user account. The system must combine a recognizable college map, location descriptions, panoramic imagery, narration and navigation graphics in a single web interface.")
add_para(doc, "A plain digital map is insufficient because it shows only position. The project therefore connects spatial selection with multimedia explanation. It also provides an administrator-controlled publishing workflow so institute content is curated rather than crowd-sourced.")

h = doc.add_heading("2. Proposed Solution", level=1); paragraph_border_bottom(h)
add_para(doc, "SGSITS Virtual Campus uses the supplied campus raster map as the primary coordinate surface. The administrator clicks a point and the browser converts the click into normalized mapX and mapY percentages. These percentages stay aligned when the map is resized. The administrator then publishes a place name, campus zone, category, visitor description, history, audio script, visiting note, accessibility information, panorama and optional gallery photographs.")
add_para(doc, "A spectator can search the place directory or click a numbered marker. Selection opens a detailed tour stop and animates a route from the Main Gate. No spectator authentication is required. This separation keeps public exploration simple while restricting content modification to the college administrator.")

h = doc.add_heading("3. Objectives", level=1); paragraph_border_bottom(h)
add_bullets(doc, [
    "Create an institute-specific virtual tour for SGSITS Indore.",
    "Allow visitors to choose a place from either the map or a sliding directory.",
    "Provide an interactive 360 degree panorama with drag and zoom controls.",
    "Give the college administrator a secure map-point publishing workflow.",
    "Store responsive normalized campus coordinates instead of fixed screen pixels.",
    "Demonstrate DDA rasterization, Bezier animation and 2D transformations in a practical navigation task.",
    "Provide gallery, history, accessibility and browser-based audio narration.",
    "Deploy the complete system using Vercel, Render and Supabase."
])

h = doc.add_heading("4. Scope of the Project", level=1); paragraph_border_bottom(h)
add_table(doc, ["In Scope", "Out of Scope / Future Work"], [
    ("Public campus exploration without login", "Indoor floor-by-floor navigation"),
    ("Single college administrator authentication", "Multiple departmental editor roles"),
    ("Map-based stop creation and selection", "GPS-based turn-by-turn mobile navigation"),
    ("Panorama and gallery upload", "Dedicated object storage and CDN pipeline"),
    ("Browser speech synthesis narration", "Professionally recorded multilingual audio"),
    ("Canvas-based CG route demonstration", "Full WebGL digital-twin reconstruction"),
], [3.38, 3.39])

h = doc.add_heading("5. Technology Stack", level=1); paragraph_border_bottom(h)
add_table(doc, ["Layer", "Technology", "Purpose"], [
    ("Frontend", "React 19 + Vite", "Single-page responsive user interface"),
    ("Graphics", "HTML Canvas", "DDA route pixels and animated direction marker"),
    ("Media", "CSS background transforms", "Interactive panorama drag and zoom"),
    ("Backend", "Node.js + Express", "REST APIs, authentication and uploads"),
    ("Validation", "Zod + Multer", "Structured fields and image upload checks"),
    ("ORM", "Prisma 6", "Database models and PostgreSQL/SQLite access"),
    ("Local DB", "SQLite", "Simple offline development database"),
    ("Cloud DB", "Supabase PostgreSQL", "Persistent production data"),
    ("Security", "JWT + bcrypt", "Admin session and password hashing"),
    ("Hosting", "Vercel + Render", "Frontend and backend deployment"),
], [1.35, 2.05, 3.37])

h = doc.add_heading("6. System Architecture", level=1); paragraph_border_bottom(h)
add_para(doc, "The system follows a three-tier web architecture. The Vercel-hosted React client requests public tour data and submits administrator actions to the Render-hosted Express API. The API validates JWT authorization and uploaded fields, then reads or writes through Prisma to Supabase PostgreSQL. Local development uses the same application models with an explicit SQLite schema.")
add_table(doc, ["Component", "Responsibility"], [
    ("React frontend", "Campus map, directory, detail tabs, panorama, narration and admin modal"),
    ("Canvas graphics layer", "Responsive route rasterization and animated movement"),
    ("Express API", "Public reads, admin login, protected publishing and media handling"),
    ("Prisma ORM", "Typed database access with separate local and production schemas"),
    ("Supabase PostgreSQL", "Cloud persistence for administrator, places and media"),
    ("Vercel", "Static frontend hosting and VITE_API_URL configuration"),
    ("Render", "Backend runtime, database migration and seed workflow"),
], [1.75, 5.02])
add_callout(doc, "Deployment flow", "Vercel frontend -> Render REST API -> Supabase PostgreSQL")

h = doc.add_heading("7. User Roles and Main Flow", level=1); paragraph_border_bottom(h)
doc.add_heading("7.1 Spectator", level=2)
add_numbers(doc, [
    "Open the public website without creating an account.",
    "Choose a campus landmark from the image map or sliding directory.",
    "Follow the animated route from the Main Gate to the selected marker.",
    "Use 360 View, About, Gallery, Audio Guide and CG Navigation tabs.",
    "Search another campus place and continue the virtual visit."
])
doc.add_heading("7.2 College Administrator", level=2)
add_numbers(doc, [
    "Sign in with the protected administrator account.",
    "Open Add stop and click the exact location on the campus map.",
    "Enter place, zone, category, visitor description, history and access details.",
    "Upload an equirectangular panorama and optional gallery photographs.",
    "Publish the stop; the API saves normalized coordinates and media.",
    "Confirm that the new numbered marker appears for public spectators."
])

h = doc.add_heading("8. Major Application Modules", level=1); paragraph_border_bottom(h)
add_table(doc, ["Module", "Functionality"], [
    ("Landing page", "SGSITS identity, project introduction and direct tour actions"),
    ("Campus directory", "Searchable sliding list of buildings, landmarks and facilities"),
    ("Campus image map", "Numbered responsive markers and direct place selection"),
    ("Administrator login", "Single role-restricted JWT authentication flow"),
    ("Admin map studio", "Click-to-pin publishing form with panorama and gallery upload"),
    ("360 panorama", "Pointer drag and wheel zoom over uploaded equirectangular imagery"),
    ("Place information", "Description, history, campus zone, hours and accessibility"),
    ("Audio guide", "Browser SpeechSynthesis narration from administrator script"),
    ("CG navigation", "DDA route, Bezier motion, transforms and viewport mapping"),
], [1.75, 5.02])

h = doc.add_heading("9. Database Design", level=1); paragraph_border_bottom(h)
add_para(doc, "The inherited database contains User, Place, PlaceMedia and Comment models. The revised application actively uses User, Place and PlaceMedia. Comment is retained for schema compatibility but is not exposed in the spectator interface.")
add_table(doc, ["Model", "Important Fields", "Purpose"], [
    ("User", "id, name, email, passwordHash, role", "Stores the single ADMIN identity and hashed credentials"),
    ("Place", "title, district, category, story, history, mapX, mapY, panoramaDataUrl", "Stores a campus stop and its normalized map point"),
    ("PlaceMedia", "type, caption, dataUrl, mimeType, sizeBytes, placeId", "Stores optional gallery images linked to a stop"),
    ("Comment", "body, placeId, userId, createdAt", "Legacy-compatible model; not used by public tour flow"),
], [1.1, 3.25, 2.42])
add_callout(doc, "Media storage note", "Panoramas and gallery images are stored as data URLs for a compact college demonstration. Production scale should move binary assets to Supabase Storage or another object-storage service.")

h = doc.add_heading("10. API Design", level=1); paragraph_border_bottom(h)
add_table(doc, ["Method", "Endpoint", "Purpose / Access"], [
    ("GET", "/api/health", "Backend health check - public"),
    ("POST", "/api/auth/login", "Administrator login and JWT issue"),
    ("GET", "/api/me", "Current authenticated administrator"),
    ("GET", "/api/places", "List campus stops - public"),
    ("GET", "/api/places/:id", "Fetch one tour stop - public"),
    ("GET", "/api/places/:id/guide", "Structured multimedia guide - public"),
    ("POST", "/api/places", "Publish stop and images - ADMIN only"),
], [0.8, 2.15, 3.82])

h = doc.add_heading("11. Backend Processing Logic", level=1); paragraph_border_bottom(h)
add_para(doc, "The backend validates login fields, campus-stop metadata and file uploads. requireAuth verifies the JWT and loads the database user, while requireAdmin rejects every non-ADMIN account. The public endpoints return sanitized user and place data. The upload route converts image buffers to data URLs, creates PlaceMedia records, and stores the selected map point.")
add_para(doc, "CORS normalizes trailing slashes in CLIENT_URL and supports comma-separated allowed origins. Production Prisma uses PostgreSQL by default; local commands explicitly select schema.sqlite.prisma. The Render build workflow installs backend dependencies, generates the PostgreSQL client, pushes the schema and runs an idempotent seed.")
add_code(doc, [
    "mapX = ((pointerX - mapLeft) / mapWidth) * 100",
    "mapY = ((pointerY - mapTop) / mapHeight) * 100",
    "screenX = (mapX / 100) * currentMapWidth",
    "screenY = (mapY / 100) * currentMapHeight",
])

h = doc.add_heading("12. Computer Graphics Implementation", level=1); paragraph_border_bottom(h)
add_para(doc, "The CG Navigation tab explains and displays the graphics pipeline using the actual selected campus stop. The algorithms are visible in the application and directly connected to navigation behavior.")
add_table(doc, ["CG Concept", "Implementation in SGSITS Virtual Campus"], [
    ("Viewport mapping", "Normalized campus percentages are mapped to the current responsive image dimensions."),
    ("DDA line drawing", "Discrete gold pixels are plotted from the Main Gate to the chosen map point."),
    ("Quadratic Bezier curve", "The navigation arrow follows a smooth curved path using parameter t."),
    ("Translation", "Canvas context is translated to the arrow's current animated point."),
    ("Rotation", "The arrow is rotated using atan2 so it faces the route direction."),
    ("Animation loop", "requestAnimationFrame clears and redraws the Canvas continuously."),
    ("Image transformation", "Panorama background position and scale change during drag and wheel input."),
], [1.72, 5.05])
doc.add_heading("12.1 DDA Route Rasterization", level=2)
add_para(doc, "For start point (x1, y1) and destination (x2, y2), the algorithm computes dx and dy. The number of steps equals max(abs(dx), abs(dy)). It then increments x and y by dx/steps and dy/steps. Every few calculated points are drawn as small gold rectangles, producing a dotted route that retains the character of raster graphics.")
doc.add_heading("12.2 Bezier Navigation Animation", level=2)
add_code(doc, [
    "B(t) = (1 - t)^2 P0 + 2(1 - t)t P1 + t^2 P2",
    "0 <= t <= 1",
    "P0 = Main Gate, P1 = raised control point, P2 = selected stop",
])
add_para(doc, "The time parameter repeats approximately every 2.6 seconds. Translation positions the arrow at B(t), and rotation aligns it with the direction between origin and destination.")

add_page_break(doc)
h = doc.add_heading("13. Frontend UI/UX and Application Screens", level=1); paragraph_border_bottom(h)
add_para(doc, "The frontend adopts the SGSITS website identity: navy and gold colors, institute emblem, formal serif headings and clear administrative controls. The screenshots below document the implemented solution rather than the earlier MP tourism design.")
add_figure(doc, 1, "figure-1.png", "SGSITS Virtual Campus landing page with institute branding, tour actions and CG feature indicators.")

add_page_break(doc)
doc.add_heading("13.1 Administrator Authentication", level=2)
add_para(doc, "The public spectator experience does not require login. The restricted modal exists only for the college administrator and clearly explains that the account is used to publish campus stops.")
add_figure(doc, 2, "figure-2.png", "Restricted administrator login modal for protected campus-content management.")

add_page_break(doc)
doc.add_heading("13.2 Admin Map Studio", level=2)
add_para(doc, "The publishing interface combines map selection and content entry. A red draft marker records the clicked normalized coordinate. Existing numbered points remain visible to help the administrator position a new landmark relative to current tour stops.")
add_figure(doc, 3, "figure-3.png", "Administrator selecting a location and entering place, zone, category and narrative details.")

add_page_break(doc)
add_para(doc, "The lower portion of the same workflow captures visiting information, accessibility, the required panorama and optional gallery photographs. Publishing sends one multipart request to the protected backend endpoint.")
add_figure(doc, 4, "figure-4.png", "Panorama/gallery upload and Publish tour stop action in the admin workflow.")

add_page_break(doc)
doc.add_heading("13.3 Multimedia Tour Stop", level=2)
add_para(doc, "Each selected location opens a detail workspace with consistent tabs. The panorama viewer supports pointer dragging and wheel zoom. The screenshot uses a demonstration equirectangular panorama; campus-specific panorama images can be uploaded by the administrator using the same workflow.")
add_figure(doc, 5, "figure-5.png", "Interactive 360 degree tour-stop view with draggable panoramic media.")

add_page_break(doc)
doc.add_heading("13.4 Accessible Audio Narration", level=2)
add_para(doc, "The Audio Guide tab uses the browser SpeechSynthesis interface. The administrator's script is narrated with play and stop controls, providing hands-free access to place information without storing a separate audio file.")
add_figure(doc, 6, "figure-6.png", "Audio Guide tab narrating the published campus-stop script.")

add_page_break(doc)
doc.add_heading("13.5 CG Navigation View", level=2)
add_para(doc, "The CG view overlays the campus raster map with numbered points and the gold DDA route. An animated direction marker moves from the Main Gate to the currently selected location. This screen provides direct evidence of the assignment-specific graphics implementation.")
add_figure(doc, 7, "figure-7.png", "CG Navigation tab showing normalized markers and the DDA route from the Main Gate.")

add_page_break(doc)
h = doc.add_heading("14. Local Setup Guide", level=1); paragraph_border_bottom(h)
add_para(doc, "Run the following commands from the repository root in PowerShell or a terminal with Node.js installed.")
add_code(doc, [
    "npm install",
    "npm run install:all",
    "npm run db:generate",
    "npm run db:init",
    "npm run db:seed",
    "npm run dev",
])
add_table(doc, ["Service", "Local URL / Credential"], [
    ("Frontend", "http://localhost:5173"),
    ("Backend health", "http://localhost:4000/api/health"),
    ("Admin email", "admin@sgsits.ac.in"),
    ("Demo password", "admin123"),
], [1.75, 5.02])
add_callout(doc, "Security note", "The demo credentials and local JWT secret must be changed before a real institutional production launch.")

h = doc.add_heading("15. Deployment Architecture", level=1); paragraph_border_bottom(h)
add_table(doc, ["Service", "Used For", "Important Configuration"], [
    ("Supabase", "PostgreSQL database", "DATABASE_URL in Render; use a PostgreSQL connection string"),
    ("Render", "Express API", "Build: npm run render:build; Start: npm start; set CLIENT_URL"),
    ("Vercel", "React/Vite frontend", "Set VITE_API_URL to the Render URL ending in /api"),
], [1.1, 1.75, 3.92])
doc.add_heading("15.1 Production Environment Variables", level=2)
add_code(doc, [
    "Render: DATABASE_URL=<Supabase PostgreSQL URL>",
    "Render: JWT_SECRET=<long random value>",
    "Render: CLIENT_URL=https://mp-darshan-360.vercel.app",
    "Vercel: VITE_API_URL=https://mp-darshan-360.onrender.com/api",
])
add_para(doc, "The backend normalizes a trailing slash in CLIENT_URL to prevent browser CORS mismatch. For multiple Vercel preview origins, CLIENT_URL accepts a comma-separated list.")

h = doc.add_heading("16. Testing and Validation", level=1); paragraph_border_bottom(h)
add_bullets(doc, [
    "Frontend production build completes successfully with Vite.",
    "PostgreSQL and SQLite Prisma schemas validate independently.",
    "Health endpoint returns SGSITS Virtual Campus API status.",
    "Public places endpoint returns seeded campus stops without authentication.",
    "Admin login returns a JWT and ADMIN role for valid credentials.",
    "Unauthenticated place publishing returns HTTP 401.",
    "CORS preflight returns the exact deployed Vercel origin.",
    "Map markers remain aligned at multiple viewport sizes.",
    "Panorama supports drag and zoom interactions.",
    "Audio guide starts and stops narration.",
    "DDA route and Bezier direction marker animate for selected stops.",
    "Admin upload validates an image panorama and required descriptive fields."
])

h = doc.add_heading("17. Security Considerations", level=1); paragraph_border_bottom(h)
add_bullets(doc, [
    "Administrator passwords are hashed with bcrypt.",
    "JWT protects campus-stop publishing and administrator identity endpoints.",
    "requireAdmin enforces the ADMIN role after authentication.",
    "Public visitors receive read-only access and require no session.",
    "DATABASE_URL and JWT_SECRET remain backend environment variables.",
    "VITE_API_URL is intentionally public because it is embedded in the browser bundle.",
    "Multer limits uploaded files and Zod validates structured form values.",
    "CORS accepts only configured frontend origins."
])

h = doc.add_heading("18. Limitations", level=1); paragraph_border_bottom(h)
add_bullets(doc, [
    "Panorama and gallery images stored as data URLs increase database size.",
    "The viewer is a lightweight image-space panorama rather than a WebGL sphere.",
    "Routes are visual CG demonstrations and do not follow exact walkable corridors.",
    "The supplied map is a raster screenshot, so clarity is limited at extreme zoom.",
    "There is one administrator account and no departmental approval workflow.",
    "Render's free service may sleep and delay the first API request.",
    "The demonstration panorama should be replaced by actual campus photography for final presentation."
])

h = doc.add_heading("19. Future Scope", level=1); paragraph_border_bottom(h)
add_bullets(doc, [
    "Capture verified equirectangular panoramas for every major SGSITS landmark.",
    "Move images to Supabase Storage with CDN delivery and database URLs.",
    "Create graph-based routes that follow campus roads and pedestrian paths.",
    "Add indoor floor plans and room-level navigation.",
    "Support Hindi and English recorded narration.",
    "Add multiple editor roles with review, draft and publish states.",
    "Introduce WebGL or Three.js panorama rendering and device-orientation controls.",
    "Provide accessibility-first routes and keyboard navigation.",
    "Add QR codes at physical campus points to open the corresponding virtual stop."
])

add_page_break(doc)
h = doc.add_heading("20. Conclusion", level=1); paragraph_border_bottom(h)
add_para(doc, "SGSITS Virtual Campus successfully redefines the original MP tourism project as an institution-focused tour application. It provides a simple public spectator experience and a controlled administrator publishing workflow. Visitors can select campus places from either a map or directory and learn through panorama, text, gallery, narration and route graphics.")
add_para(doc, "The project is especially suitable for a Computer Graphics submission because viewport mapping, DDA rasterization, Bezier motion, translation, rotation and continuous Canvas rendering are visible parts of a useful navigation system. At the same time, the full-stack architecture demonstrates frontend development, protected APIs, database design, file uploads and multi-service cloud deployment.")

h = doc.add_heading("21. References", level=1); paragraph_border_bottom(h)
refs = [
    "Project repository: https://github.com/ojasvagr123/mp-darshan-360",
    "Live frontend: https://mp-darshan-360.vercel.app/",
    "React documentation: https://react.dev/",
    "Vite documentation: https://vite.dev/",
    "MDN Canvas API: https://developer.mozilla.org/docs/Web/API/Canvas_API",
    "Express documentation: https://expressjs.com/",
    "Prisma documentation: https://www.prisma.io/docs/",
    "Supabase documentation: https://supabase.com/docs",
    "Render documentation: https://render.com/docs",
    "Vercel documentation: https://vercel.com/docs",
]
add_numbers(doc, refs)

# Core properties
doc.core_properties.title = "SGSITS Virtual Campus - Project Report"
doc.core_properties.subject = "Computer Graphics project report"
doc.core_properties.author = "Ojasv Agrawal"
doc.core_properties.keywords = "SGSITS, virtual campus, computer graphics, DDA, Bezier, panorama"
doc.core_properties.comments = "Updated from the earlier MP Darshan concept for the SGSITS campus-tour solution."

doc.save(OUT)
print(OUT)
